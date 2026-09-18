# Copyright 2025-2026 Hanlian Lu. SPDX-License-Identifier: Apache-2.0
"""Round-one resource regressions through real Host settlement and disposable PG."""

from __future__ import annotations

import asyncio
import base64
import io
import json
import threading
import uuid
import zipfile
from contextlib import contextmanager
from dataclasses import asdict, replace
from types import SimpleNamespace
from typing import Any, cast
from unittest.mock import AsyncMock, patch

import pytest
from docx import Document

from dlightrag.adapters.postgres.runtime.run_blob_store import PGRunBlobStore
from dlightrag.engine.agent.environment.confinement import ConfinementPolicy
from dlightrag.engine.agent.session.fold import PriorTurns, project_session_messages
from dlightrag.engine.agent.session.ids import LaneId
from dlightrag.engine.ai.capacity import CONTEXT_POLICY_REVISION
from dlightrag.engine.ai.catalog import current_model_catalog_revision
from dlightrag.engine.ai.fingerprints import ModelFingerprint
from dlightrag.engine.ai.messages import AssistantTurn, ToolCall
from dlightrag.engine.ai.providers import get_provider as real_get_provider
from dlightrag.engine.ai.scheduler import ModelScheduler
from dlightrag.engine.ai.settings import (
    CHAT_MODEL_SELECTORS,
    ModelRoleOverrides,
    ModelRoleSettings,
    ModelSettings,
)
from dlightrag.engine.ai.telemetry import NOOP_TELEMETRY
from dlightrag.engine.answer.capabilities import RequestModelContext
from dlightrag.engine.answer.errors import AnswerInputOverflowError
from dlightrag.engine.answer.execution.executor import (
    AnswerExecutor,
    AnswerExecutorSettings,
    AnswerResourceResolver,
    AnswerResourceSettings,
)
from dlightrag.engine.answer.execution.input import (
    AnswerRunInput,
    PinnedModelProfile,
    model_reasoning_settings,
)
from dlightrag.engine.answer.fast import ensure_session_lane
from dlightrag.engine.answer.highlights import SemanticHighlightSettings
from dlightrag.engine.answer.history import HistoryInputMeasure
from dlightrag.engine.answer.images import AnswerImageBudget, AnswerImagePolicy
from dlightrag.engine.answer.model_runtime import (
    AnswerModelRuntime,
    AnswerModelRuntimeSettings,
    WebSourceRuntimeSettings,
)
from dlightrag.engine.answer.orchestration import AnswerOrchestrator
from dlightrag.engine.answer.research.runtime import (
    FetchedResourceBuffer,
    _bound_child_dispatch_preparer,
    _bound_child_runner,
    _fenced_child_writer,
)
from dlightrag.engine.answer.resources.models import (
    ResourceAdmissionError,
    ResourceInput,
    TextWindowBudget,
)
from dlightrag.engine.answer.resources.registry import ResourceRegistry
from dlightrag.engine.answer.runs.routing import RoutingAcceptance
from dlightrag.engine.answer.synthesizer import AnswerSynthesizer
from dlightrag.engine.answer.tools.subagents import SubagentHost
from dlightrag.engine.rag.retrieval import RetrievalResult
from dlightrag.engine.runtime.coordinator import RunCoordinator
from tests.integration.run_runtime_pg_harness import run_envelope
from tests.integration.test_attachment_replay_pg import (  # noqa: F401
    OWNER,
    drive,
    executor,
    new_run,
    orchestrator,
    origin,
)
from tests.unit.conftest import answer_image_policy, answer_model_profile
from tests.unit.test_answer_executor import _resource_resolver
from tests.unit.test_docx_conversion import _zip_replace
from tests.unit.test_provider_attachment_contract import (
    _openai_complete_json,
    _openai_stream_sse,
    bind_mock_http,
)
from tests.unit.test_resource_tools import png as png_bytes

pytestmark = [pytest.mark.integration, pytest.mark.asyncio]


@pytest.fixture
async def pg():
    from tests.integration.run_runtime_pg_harness import isolated_run_runtime

    async with isolated_run_runtime("resource_review") as pair:
        async with pair[1].acquire() as conn:
            assert int(await conn.fetchval("SHOW server_version_num")) >= 180000
        yield pair


def refused_docx(terminal):
    output = io.BytesIO()
    if terminal == "archive":
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", b"x" * 1_000_000)
        return output.getvalue()
    document = Document()
    document.add_paragraph("Generated depth refusal")
    document.save(output)
    deep = (
        b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        + b"<w:body>" * 260
        + b"</w:body>" * 260
        + b"</w:document>"
    )
    return _zip_replace(output.getvalue(), {"word/document.xml": deep})


@pytest.mark.parametrize("admission", ["agent", "caller"])
@pytest.mark.parametrize("terminal", ["archive", "native_limit"])
async def test_url_terminal_source_and_conversion_settle_together(
    pg, monkeypatch, admission, terminal
):
    session, session_id = await new_run(pg[0])
    data = refused_docx(terminal)
    url = "https://example.com/generated-refusal.docx"
    fetch = AsyncMock(
        return_value=SimpleNamespace(
            content=data,
            final_url=url,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    )
    monkeypatch.setattr("dlightrag.engine.answer.resources.registry.fetch_public_http", fetch)
    monkeypatch.setattr(
        "dlightrag.engine.answer.resources.converters._convert_markitdown",
        lambda *a, **k: pytest.fail("terminal must not fall back"),
    )
    buffer = FetchedResourceBuffer()

    async def sink(fetched, owner):
        buffer.append(fetched, owner)

    async with ResourceRegistry(
        resource_secret=b"url-refusal", fetched_bytes_sink=sink
    ) as registry:
        resource = registry.register(ResourceInput(url=url)) if admission == "caller" else None
        calls = 0

        async def model(**kwargs):
            nonlocal calls
            calls += 1
            if calls == 1:
                return AssistantTurn(
                    text="",
                    tool_calls=(
                        ToolCall(
                            "read", "read", {"resource_id": resource} if resource else {"url": url}
                        ),
                    ),
                    stop_reason="tool_use",
                )
            assert "safety_refused" in str(kwargs["messages"])
            return AssistantTurn(text="refused honestly", tool_calls=(), stop_reason="stop")

        host = orchestrator(model, registry=registry)
        await drive(
            session,
            session_id,
            host,
            host.prepare_run("read refusal", registry=registry),
            fetched_buffer=buffer,
        )
        resource = registry.manifest()[0].resource_id
        effects = registry.conversion_effects(resource)
        assert effects and fetch.await_count == 1

    async def forbidden(*args, **kwargs):
        raise AssertionError("settled refusal must not fetch or parse")

    monkeypatch.setattr("dlightrag.engine.answer.resources.registry.fetch_public_http", forbidden)
    monkeypatch.setattr("dlightrag.engine.answer.resources.registry.convert_resource", forbidden)
    async with ResourceRegistry(resource_secret=b"url-refusal") as restored:
        if admission == "caller":
            assert restored.register(ResourceInput(url=url)) == resource
        await executor(pg)._restore_registry_fetches(
            restored, owner_id=OWNER, run_id=session.run_id
        )
        assert restored.conversion_effects(resource) == effects
        assert (await restored.materialize(resource)) == data
        from tests.unit.test_resource_tools import call, tools

        read, view = tools(restored)
        for _ in range(2):
            result = await call(read, resource_id=resource)
            assert result.is_error and "safety_refused" in result.text_content
        with pytest.raises(ResourceAdmissionError, match="previously refused"):
            await call(view, resource_id=resource)


@pytest.mark.parametrize(
    "max_images,child_views,child_context",
    [(1, 0, 64_000), (2, 0, 64_000), (3, 1, 64_000), (2, 0, 1024)],
)
async def test_same_host_async_spawn_dispatch_retry_and_continuation_budget(
    pg, max_images, child_views, child_context
):
    session, session_id = await new_run(pg[0])
    store = pg[0]
    child_seen = []
    async with ResourceRegistry() as registry:
        resource = registry.register(ResourceInput(filename="generated.png", content=png_bytes()))
        parent_calls = 0

        async def parent_model(**kwargs):
            nonlocal parent_calls
            parent_calls += 1
            if parent_calls == 1:
                return AssistantTurn(
                    text="",
                    tool_calls=(ToolCall("view", "view", {"resource_id": resource}),),
                    stop_reason="tool_use",
                )
            if parent_calls == 2:
                return AssistantTurn(
                    text="",
                    tool_calls=(
                        ToolCall(
                            "spawn",
                            "spawn_agent",
                            {
                                "children": [
                                    {
                                        "objective": "Review inherited pixels",
                                        "context": "parent",
                                        "tools": ["view"] if child_views else [],
                                    }
                                ]
                            },
                        ),
                    ),
                    stop_reason="tool_use",
                )
            return AssistantTurn(text="parent done", tool_calls=(), stop_reason="stop")

        async def child_model(**kwargs):
            attachments = [a for m in kwargs["messages"] for a in m.get("attachments", [])]
            expected = 1 + (child_views if child_seen else 0)
            assert len(attachments) == expected and attachments[0]["data_url"].startswith(
                "data:image/"
            )
            child_seen.append(kwargs["model_profile"])
            if child_views and len(child_seen) == 1:
                return AssistantTurn(
                    text="",
                    tool_calls=(ToolCall("child-view", "view", {"resource_id": resource}),),
                    stop_reason="tool_use",
                )
            return AssistantTurn(text="child done", tool_calls=(), stop_reason="stop")

        host = orchestrator(parent_model, registry=registry, max_images=max_images)
        child_profile = replace(
            answer_model_profile(supports_images=True), context_window_tokens=child_context
        )
        host._child_model_resolver = cast(Any, lambda role: (child_model, None, child_profile))
        budget = host._image_budget
        assert budget is not None
        persist = _fenced_child_writer(store, "upsert_child_session", session)
        claim = _fenced_child_writer(store, "claim_child_session", session)
        renew = _fenced_child_writer(store, "heartbeat_child_session", session)
        assert persist and claim and renew
        dispatch = _bound_child_dispatch_preparer(host)

        def retried_dispatch(*args):
            before = budget.count
            first = dispatch(*args)
            assert dispatch(*args) == first
            assert budget.count == before
            return first

        subagents = SubagentHost(
            parent_session_id=session_id,
            owner_id=OWNER,
            run_id=session.run_id,
            persist=persist,
            load_child=store.load_child_session,
            list_children=store.list_child_sessions,
            finish_child=_fenced_child_writer(store, "finish_child_session", session),
            prepare_dispatch=retried_dispatch,
            run_child=_bound_child_runner(
                telemetry=NOOP_TELEMETRY,
                orchestrator=host,
                repository=session.execution.session_repository,
                session=session,
                fetched_buffer=FetchedResourceBuffer(),
                parent_session_id=session_id,
                persist_child_runtime=persist,
                claim_child=claim,
                renew_child=renew,
                load_child=store.load_child_session,
                restore_child_attachments=lambda child, context: executor(
                    pg
                )._restore_child_attachments(session, child, context),
            ),
        )
        host._subagent_host = subagents
        await drive(
            session, session_id, host, host.prepare_run("view and spawn", registry=registry)
        )
        assert len(subagents.tasks) == 1
        outcomes = await asyncio.wait_for(asyncio.gather(*subagents.tasks.values()), 10)
        if max_images == 1 or child_context == 1024:
            assert outcomes[0].status == "failed" and child_seen == []
            return
        assert outcomes[0].status == "succeeded" and child_seen == [child_profile] * (
            1 + child_views
        )
        assert budget.count == 2 + child_views
        child_id = outcomes[0].child_session_id
        receipt = await store.continue_child_session(
            owner_id=OWNER,
            run_id=session.run_id,
            child_session_id=child_id,
            content="Verify again",
            submission_key="generated-continuation",
        )
        assert receipt["outcome"] == "accepted"
        await subagents.restore_pending()
        outcomes = await asyncio.wait_for(asyncio.gather(*subagents.tasks.values()), 10)
        assert outcomes[0].status == "succeeded" and child_seen == [child_profile] * (
            2 + child_views
        )
        assert budget.count == 2 + child_views  # Same occurrences, not a new admission.


@pytest.mark.parametrize("kind", ["follow_up", "fork"])
@pytest.mark.parametrize(
    "own_images,max_images,vision",
    [(0, 2, True), (0, 1, True), (1, 3, True), (1, 2, True), (0, 3, False)],
)
async def test_research_view_to_fast_uses_one_consuming_budget(
    pg, kind, own_images, max_images, vision
):
    _, session_id, snapshot, _, _, _ = await origin(pg)
    lane = LaneId.main() if kind == "follow_up" else LaneId.new()
    current, _ = await new_run(
        pg[0],
        session_id=session_id,
        lane=lane.value,
        source_lane="main" if kind == "fork" else None,
        mode="fast",
    )
    await ensure_session_lane(
        repository=current.execution.session_repository,
        snapshot=snapshot,
        fencing_epoch=current.fencing_epoch,
        session_id=session_id,
        lane_id=lane,
        source_lane_id=LaneId.main() if kind == "fork" else None,
    )
    snapshot = replace(
        await current.execution.session_repository.load(session_id), selected_lane_id=lane
    )
    snapshots = await executor(pg)._restore_selected_attachments(current, snapshot)
    messages = project_session_messages(snapshot.tree.ancestry(lane), snapshot.active_projection)
    resolver = _resource_resolver()
    profile = answer_model_profile(supports_images=vision)
    policy = answer_image_policy(max_images=max_images if vision else 0)
    capabilities = cast(Any, resolver._capabilities)
    capabilities.answer_image_policy.return_value = policy
    from dlightrag.engine.answer.capabilities import RequestModelContext

    models = RequestModelContext(query=profile, extract=profile, vlm=profile)

    async def confirm(models):
        return models, resolver._capabilities.answer_capability_from_profile(profile)

    # Current-image capability is independently confirmed, as at request ingress.
    from dlightrag.engine.answer.image_capability import AnswerImageCapability

    capabilities.answer_capability_from_profile.return_value = AnswerImageCapability(
        status="supported",
        configured_ceiling=max_images,
        effective_max_images=max_images,
        provider="generated",
        base_url=None,
        model="generated",
        failure_kind=None,
    )
    resolved = await resolver.resolve(
        resources=[ResourceInput(filename="own.png", content=png_bytes())] * own_images,
        models=models,
        confirm_image_context=confirm,
        resolved_mode="fast",
    )
    assert resolved.image_budget is not None
    seen = []

    async def model(**kwargs):
        seen.append(kwargs["messages"])

        async def stream():
            yield "generated fast response"

        return stream()

    async def retrieve(*args, **kwargs):
        return RetrievalResult()

    host = AnswerOrchestrator(
        synthesizer=AnswerSynthesizer(image_policy=policy, model_profile=profile, model_func=model),
        retrieve_knowledge_base=retrieve,
        model_profile=profile,
        image_budget=resolved.image_budget,
        text_window_budget=TextWindowBudget(4000),
        telemetry=NOOP_TELEMETRY,
        resolved_mode="fast",
    )
    try:
        if not vision or own_images + 2 > max_images:
            with pytest.raises(AnswerInputOverflowError, match="image budget"):
                host.admit_durable_attachments(messages, snapshots)
            assert not seen
            return
        host.admit_durable_attachments(messages, snapshots)
        assert resolved.image_budget.count == own_images + 2
        _, stream = await host.answer_stream(
            "Follow up",
            conversation_history=PriorTurns(messages),
            query_images=resolved.query_images,
        )
        assert stream is not None
        _ = [part async for part in stream]
        assert len(seen) == 1
        assert sum(len(m.get("attachments", [])) for m in seen[0]) == 2
        assert (
            sum(
                1
                for m in seen[0]
                for b in (m.get("content") if isinstance(m.get("content"), list) else ())
                if b.get("type") == "image_url"
            )
            == own_images
        )
        assert resolved.image_budget.count == own_images + 2
    finally:
        if resolved.registry:
            await resolved.registry.aclose()


class _ObservedExecutor:
    def __init__(self, executor: AnswerExecutor) -> None:
        self.executor = executor
        self.failed = asyncio.Event()
        self.error: BaseException | None = None

    async def execute(self, session):
        try:
            return await self.executor.execute(session)
        except BaseException as exc:
            self.error = exc
            self.failed.set()
            raise


class _RecordingImagePolicy(AnswerImagePolicy):
    """Run image policy that records every created budget for no-recharge asserts."""

    __slots__ = ("_recorder",)

    def __init__(self, recorder: list[AnswerImageBudget], **fields: Any) -> None:
        super().__init__(**fields)
        object.__setattr__(self, "_recorder", recorder)

    def new_budget(self, *, max_px: int | None = None) -> AnswerImageBudget:
        budget = super().new_budget(max_px=max_px)
        self._recorder.append(budget)
        return budget


class _FastProjectionProvider:
    """HTTP recorder for Fast follow-up/fork through AnswerModelRuntime."""

    def __init__(self, *, block: bool = False) -> None:
        self.block = block
        self.started = asyncio.Event()
        self.wire_requests: list[dict[str, Any]] = []
        self.budgets: list[AnswerImageBudget] = []

    async def handler(self, request: Any) -> Any:
        import httpx2

        body = json.loads(request.content.decode())
        self.wire_requests.append(body)
        self.started.set()
        if self.block:
            await asyncio.Event().wait()
        if body.get("stream"):
            return httpx2.Response(
                200,
                headers={"content-type": "text/event-stream"},
                content=_openai_stream_sse(),
            )
        return httpx2.Response(200, json=_openai_complete_json())


@contextmanager
def _http_bound_providers(recorder: _FastProjectionProvider):
    def factory(name: str, **kwargs: Any):
        provider = real_get_provider(name, **kwargs)
        bind_mock_http(provider, recorder.handler)
        return provider

    with (
        patch("dlightrag.engine.ai.completion.get_provider", factory),
        patch("dlightrag.engine.ai.tool_model.get_provider", factory),
    ):
        yield


def _fast_role_settings() -> ModelRoleSettings:
    def one(role: str) -> ModelSettings:
        return ModelSettings(
            provider="openai",
            model=f"projection-{role}",
            api_key="test-key",
            max_retries=0,
        )

    return ModelRoleSettings(
        default=one("default"),
        roles=ModelRoleOverrides(
            extract=one("extract"),
            keyword=one("keyword"),
            query=one("query"),
            vlm=one("vlm"),
        ),
    )


def _fast_executor(pg, provider: _FastProjectionProvider, profile):
    policy = _RecordingImagePolicy(provider.budgets, **asdict(answer_image_policy(max_images=4)))
    runtime = AnswerModelRuntime(
        settings=AnswerModelRuntimeSettings(
            model_roles=_fast_role_settings(),
            web_sources=WebSourceRuntimeSettings(),
            query_image_limit=4,
        ),
        scheduler=ModelScheduler(max_concurrency=1),
        telemetry=NOOP_TELEMETRY,
        answer_image_policy=lambda _profile: policy,
        vlm_image_policy=lambda _profile: policy,
        vlm_profile=lambda: profile,
    )

    class Capabilities:
        @staticmethod
        def request_model_context(pinned):
            return RequestModelContext(
                query=pinned["query"], extract=pinned["extract"], vlm=pinned["vlm"]
            )

        @staticmethod
        def answer_image_policy(_profile):
            return policy

        @staticmethod
        async def pinned_answer_context(models):
            return models, None

    models = runtime
    capabilities = Capabilities()

    async def retrieve(*args, **kwargs):
        return RetrievalResult()

    async def planner_history_input_measure(**kwargs) -> HistoryInputMeasure:
        def measure(messages: list[dict[str, Any]], projected_summary: str = "") -> int:
            return 1

        return measure

    return AnswerExecutor(
        store=pg[0],
        blob_store=PGRunBlobStore(pool=pg[1]),
        pool=cast(Any, SimpleNamespace()),
        warm=lambda _workspaces: None,
        retrieve=retrieve,
        planner_history_input_measure=planner_history_input_measure,
        models=cast(Any, models),
        capabilities=cast(Any, capabilities),
        resources=AnswerResourceResolver(
            settings=AnswerResourceSettings(
                max_attachments=6,
                max_attachment_bytes=10_000_000,
                max_total_attachment_bytes=20_000_000,
                image_max_bytes=10_000_000,
                image_max_pixels=10_000_000,
            ),
            models=cast(Any, models),
            capabilities=cast(Any, capabilities),
        ),
        settings=AnswerExecutorSettings(
            default_top_k=10,
            default_chunk_top_k=20,
            semantic_highlights=SemanticHighlightSettings(
                enabled=False,
                timeout=1.0,
                max_concurrency=1,
                batch_size=1,
                max_input_chars=100,
                cache_size=1,
            ),
        ),
        telemetry=NOOP_TELEMETRY,
        model_fingerprint_for_role=lambda role: ModelFingerprint(
            "openai", f"projection-{role}", None
        ),
        execution_environment="disabled",
        shell_confinement=ConfinementPolicy(),
    )


async def _create_fast_continuation(pg, *, old_run_id, session_id, kind, profile):
    lane = LaneId.main() if kind == "follow_up" else LaneId.new()
    settings = {role: ModelSettings(model=f"projection-{role}") for role in CHAT_MODEL_SELECTORS}
    prepared = AnswerRunInput(
        query="Use the exact historical pages",
        workspaces=("workspace-000",),
        pinned_models=tuple(
            PinnedModelProfile(
                role=role,
                fingerprint=ModelFingerprint("openai", f"projection-{role}", None),
                profile=profile,
                reasoning_settings=model_reasoning_settings(settings[role]),
            )
            for role in CHAT_MODEL_SELECTORS
        ),
        context_policy_revision=CONTEXT_POLICY_REVISION,
        model_catalog_revision=current_model_catalog_revision(),
        idempotency_fingerprint=f"fast-projection-{uuid.uuid7()}",
        agent_session_id=session_id.value,
        agent_lane_id=lane.value,
        source_lane_id="main" if kind == "fork" else None,
        parent_run_id=old_run_id,
        continuation_kind=kind,
    )
    run_id = str(uuid.uuid7())
    envelope = replace(
        run_envelope("answer", key=run_id, owner=OWNER, mode="fast"),
        payload=prepared.as_request(),
        accepted_input={
            "query": prepared.query,
            "workspaces": list(prepared.workspaces),
            "mode": "fast",
        },
    )
    routing = RoutingAcceptance(
        requested_mode="fast",
        valid_modes=("fast",),
        resolved_mode="fast",
        context_policy_revision=CONTEXT_POLICY_REVISION,
        model_fingerprints={},
        agent_session_id=session_id.value,
        agent_lane_id=lane.value,
        source_lane_id=prepared.source_lane_id,
    )
    await pg[0].create_run(envelope=envelope, run_id=run_id, routing=routing)
    return run_id


def _assert_fast_provider_pixels(provider, expected):
    assert len(provider.wire_requests) == 1
    wire = provider.wire_requests[0]["messages"]
    for message in wire:
        assert (
            not {
                "attachments",
                "provider_state",
                "is_error",
                "untrusted_tool_data",
            }
            & message.keys()
        )
    wire_images = [
        block["image_url"]["url"]
        for message in wire
        for block in (message.get("content") if isinstance(message.get("content"), list) else ())
        if block.get("type") == "image_url"
    ]
    assert len(wire_images) == 2
    assert all(base64.b64decode(item.partition(",")[2]) == expected for item in wire_images)
    # Both occurrences are reserved exactly once; re-hydration after projection
    # must not charge the Fast run budget a second time.
    charged = [budget for budget in provider.budgets if budget.count == 2]
    assert len(charged) == 1
    assert charged[0].used_bytes == 2 * len(expected)
    assert all(budget.count <= 2 for budget in provider.budgets)


@pytest.mark.parametrize("kind", ["follow_up", "fork"])
async def test_fast_executor_rehydrates_historical_pixels_after_projection_and_restart(pg, kind):
    old, session_id, _snapshot, selection, _, _ = await origin(pg)
    digest = selection.occurrences[0].attachment.content_digest
    expected = b"".join(
        [part async for part in PGRunBlobStore(pool=pg[1]).stream(owner_id=OWNER, digest=digest)]
    )
    profile = replace(answer_model_profile(supports_images=True), context_window_tokens=1_000_000)
    run_id = await _create_fast_continuation(
        pg,
        old_run_id=old.run_id,
        session_id=session_id,
        kind=kind,
        profile=profile,
    )

    first = _FastProjectionProvider(block=True)
    with _http_bound_providers(first):
        observed = _ObservedExecutor(_fast_executor(pg, first, profile))
        coordinator = RunCoordinator(
            store=pg[0],
            executors={"answer": observed},
            query_worker_concurrency=1,
        )
        await coordinator.start()
        coordinator.wake()
        try:
            started = asyncio.create_task(first.started.wait())
            failed = asyncio.create_task(observed.failed.wait())
            done, pending = await asyncio.wait(
                (started, failed), timeout=10, return_when=asyncio.FIRST_COMPLETED
            )
            for task in pending:
                task.cancel()
            await asyncio.gather(*pending, return_exceptions=True)
            assert done, "Fast continuation neither called its provider nor reported failure"
            if observed.error is not None:
                raise observed.error
        finally:
            await coordinator.aclose()
    _assert_fast_provider_pixels(first, expected)

    queued = await pg[0].get_run(owner_id=OWNER, run_id=run_id)
    assert queued is not None and queued.status == "queued"
    resumed = _FastProjectionProvider()
    with _http_bound_providers(resumed):
        resumed_observed = _ObservedExecutor(_fast_executor(pg, resumed, profile))
        coordinator = RunCoordinator(
            store=pg[0],
            executors={"answer": resumed_observed},
            query_worker_concurrency=1,
        )
        await coordinator.start()
        coordinator.wake()
        try:
            for _ in range(500):
                record = await pg[0].get_run(owner_id=OWNER, run_id=run_id)
                if record is not None and record.status == "succeeded":
                    break
                if resumed_observed.error is not None:
                    raise resumed_observed.error
                await asyncio.sleep(0.02)
            else:
                raise AssertionError("resumed Fast continuation did not settle")
        finally:
            await coordinator.aclose()
    _assert_fast_provider_pixels(resumed, expected)


async def test_cancelled_child_url_terminal_settles_source_for_parent_and_recovery(pg, monkeypatch):
    data_stream = io.BytesIO()
    document = Document()
    document.add_paragraph("Native cancellation fixture")
    document.save(data_stream)
    data = data_stream.getvalue()
    url = "https://example.com/cancelled-child.docx"
    native_started = threading.Event()
    native_release = threading.Event()

    def native(*args, **kwargs):
        native_started.set()
        assert native_release.wait(5)
        return "late output must not be adopted"

    fetch = AsyncMock(
        return_value=SimpleNamespace(
            content=data,
            final_url=url,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    )
    monkeypatch.setattr("anydoc.to_markdown_bytes", native)
    monkeypatch.setattr("dlightrag.engine.answer.resources.registry.fetch_public_http", fetch)
    monkeypatch.setattr(
        "dlightrag.engine.answer.resources.converters._convert_markitdown",
        lambda *a, **k: pytest.fail("cancelled conversion must not fall back"),
    )
    session, session_id = await new_run(pg[0])
    buffer = FetchedResourceBuffer()
    sink_owners = []

    async def sink(fetched, owner):
        sink_owners.append((fetched.resource_id, owner))
        buffer.append(fetched, owner)

    async with ResourceRegistry(
        resource_secret=b"cancelled-url", fetched_bytes_sink=sink
    ) as registry:
        parent_calls = 0
        child_calls = 0

        async def parent_model(**kwargs):
            nonlocal parent_calls
            parent_calls += 1
            if parent_calls == 1:
                return AssistantTurn(
                    text="",
                    tool_calls=(
                        ToolCall(
                            "spawn-reader",
                            "spawn_agent",
                            {
                                "children": [
                                    {
                                        "objective": "Read the URL",
                                        "context": "isolated",
                                        "tools": ["read"],
                                    }
                                ]
                            },
                        ),
                    ),
                    stop_reason="tool_use",
                )
            if parent_calls == 2:
                for _ in range(500):
                    if native_started.is_set():
                        break
                    await asyncio.sleep(0.01)
                else:
                    raise AssertionError("child conversion did not start")
                children = await pg[0].list_child_sessions(owner_id=OWNER, run_id=session.run_id)
                assert len(children) == 1
                child_id = children[0]["child_session_id"]
                asyncio.get_running_loop().call_later(0.05, native_release.set)
                return AssistantTurn(
                    text="",
                    tool_calls=(
                        ToolCall(
                            "cancel-reader",
                            "cancel_subagent",
                            {"child_session_id": child_id},
                        ),
                    ),
                    stop_reason="tool_use",
                )
            if parent_calls == 3:
                for _ in range(500):
                    manifest = registry.manifest()
                    if manifest and registry.conversion_effects(manifest[0].resource_id):
                        break
                    await asyncio.sleep(0.01)
                else:
                    raise AssertionError("cancelled conversion did not latch its terminal")
                return AssistantTurn(
                    text="",
                    tool_calls=(ToolCall("parent-read", "read", {"url": url}),),
                    stop_reason="tool_use",
                )
            assert "safety_refused" in str(kwargs["messages"])
            return AssistantTurn(text="parent retained refusal", tool_calls=(), stop_reason="stop")

        async def child_model(**kwargs):
            nonlocal child_calls
            child_calls += 1
            if child_calls > 1:
                raise AssertionError("cancelled child must not reach another model turn")
            return AssistantTurn(
                text="",
                tool_calls=(ToolCall("child-read", "read", {"url": url}),),
                stop_reason="tool_use",
            )

        host = orchestrator(parent_model, registry=registry)
        host._child_model_resolver = cast(
            Any,
            lambda role: (
                child_model,
                None,
                answer_model_profile(supports_images=True),
            ),
        )
        persist = _fenced_child_writer(pg[0], "upsert_child_session", session)
        claim = _fenced_child_writer(pg[0], "claim_child_session", session)
        renew = _fenced_child_writer(pg[0], "heartbeat_child_session", session)
        assert persist and claim and renew
        subagents = SubagentHost(
            parent_session_id=session_id,
            owner_id=OWNER,
            run_id=session.run_id,
            persist=persist,
            load_child=pg[0].load_child_session,
            list_children=pg[0].list_child_sessions,
            finish_child=_fenced_child_writer(pg[0], "finish_child_session", session),
            request_cancel=_fenced_child_writer(pg[0], "request_child_cancellation", session),
            release_children=_fenced_child_writer(pg[0], "release_child_sessions", session),
            prepare_dispatch=_bound_child_dispatch_preparer(host),
            run_child=_bound_child_runner(
                telemetry=NOOP_TELEMETRY,
                orchestrator=host,
                repository=session.execution.session_repository,
                session=session,
                fetched_buffer=buffer,
                parent_session_id=session_id,
                persist_child_runtime=persist,
                claim_child=claim,
                renew_child=renew,
                load_child=pg[0].load_child_session,
                restore_child_attachments=lambda child, context: executor(
                    pg
                )._restore_child_attachments(session, child, context),
            ),
        )
        host._subagent_host = subagents
        await drive(
            session,
            session_id,
            host,
            host.prepare_run("cancel child then read", registry=registry),
            fetched_buffer=buffer,
        )
        resource = registry.manifest()[0].resource_id
        effects = registry.conversion_effects(resource)
        assert effects and fetch.await_count == 1 and child_calls == 1
        assert [item[0] for item in sink_owners] == [resource]
        durable_resources = await pg[0].list_fetched_resources(
            owner_id=OWNER, run_id=session.run_id
        )
        assert [
            (item.resource_id, item.capabilities.get("resource_kind")) for item in durable_resources
        ] == [
            (resource, "web"),
            (f"{resource}-conversion", "conversion_snapshot"),
        ]
        rows = await pg[0].list_child_sessions(owner_id=OWNER, run_id=session.run_id)
        assert len(rows) == 1 and rows[0]["status"] == "cancelled"

    async def forbidden(*args, **kwargs):
        raise AssertionError("durable cancellation terminal must not fetch or parse")

    monkeypatch.setattr("dlightrag.engine.answer.resources.registry.fetch_public_http", forbidden)
    monkeypatch.setattr("dlightrag.engine.answer.resources.registry.convert_resource", forbidden)
    async with ResourceRegistry(resource_secret=b"cancelled-url") as restored:
        snapshots = await executor(pg)._restore_registry_fetches(
            restored, owner_id=OWNER, run_id=session.run_id
        )
        assert snapshots[resource] == data
        assert await restored.materialize(resource) == data
        from tests.unit.test_resource_tools import call, tools

        read, _ = tools(restored)
        result = await call(read, resource_id=resource)
        assert result.is_error and "safety_refused" in result.text_content
        assert restored.conversion_effects(resource) == effects
