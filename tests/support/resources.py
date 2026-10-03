# Copyright 2025-2026 Hanlian Lu. SPDX-License-Identifier: Apache-2.0
"""Resource bytes, and the read and view tools bound to one registry.

Production composes read and view through the answer tool registry; tests that
exercise them over a registry they built assemble them here.
"""

import base64
import io
import re
from typing import Any

from docx import Document
from PIL import Image

from dlightrag.engine.agent.environment import AccessScheduler
from dlightrag.engine.agent.tools import AgentTool, ToolResult
from dlightrag.engine.agent.tools.files import PreparedImageAttachment, read_tool, view_tool
from dlightrag.engine.ai.media import decode_image_base64
from dlightrag.engine.answer.tools.resources import make_resource_reader, make_resource_viewer
from tests.tool_helpers import tool_runtime
from tests.unit.conftest import answer_image_policy


def png() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (24, 24), (20, 10, 0)).save(buffer, "PNG")
    return buffer.getvalue()


def pdf_bytes(count: int = 3, size: tuple[int, int] = (300, 400)) -> bytes:
    images = [Image.new("RGB", size, (index * 20, 10, 10)) for index in range(count)]
    buffer = io.BytesIO()
    images[0].save(buffer, "PDF", save_all=True, append_images=images[1:])
    for image in images:
        image.close()
    return buffer.getvalue()


def docx_images(count: int) -> bytes:
    doc = Document()
    doc.add_paragraph("Revenue was 123.")
    for _ in range(count):
        doc.add_picture(io.BytesIO(png()))
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def preparer(max_images: int = 8):
    budget = answer_image_policy(max_images=max_images).new_budget()

    def prepare(data: bytes, label: str) -> PreparedImageAttachment | None:
        block = budget.add_base64(base64.b64encode(data).decode(), label=label)
        if block is None:
            return None
        content, media = decode_image_base64(block["image_url"]["url"])
        return PreparedImageAttachment(content, media or "image/png", content != data)

    return prepare


def tools(
    registry: Any, *, max_images: int = 8, environment: Any = None
) -> tuple[AgentTool, AgentTool]:
    access = AccessScheduler()
    return (
        read_tool(
            environment,
            access,
            resource_reader=make_resource_reader(registry, 1000),
        ),
        view_tool(
            environment,
            access,
            resource_viewer=make_resource_viewer(registry),
            image_preparer=preparer(max_images),
        ),
    )


async def call(tool: AgentTool, **args: Any) -> ToolResult:
    return await tool.execute(
        tool.input_model.model_validate(args), tool_runtime(tool_name=tool.name)
    )


def printed_handle(result: ToolResult) -> str:
    """The Resource handle a read printed, which the model names from then on."""
    printed = re.search(r"\[resource: (res-[0-9a-f]+)", result.text_content)
    assert printed is not None, result.text_content
    return printed.group(1)
