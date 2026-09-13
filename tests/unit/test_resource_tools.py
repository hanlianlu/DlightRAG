# Copyright 2025-2026 Hanlian Lu. SPDX-License-Identifier: Apache-2.0
"""Public read/view seams: text, pixels, inventories, identity, and budgets."""

import base64
import io

import pytest
from docx import Document
from PIL import Image
from pydantic import ValidationError

from dlightrag.engine.agent.environment import AccessScheduler
from dlightrag.engine.agent.environment.local import LocalExecutionEnvironment
from dlightrag.engine.agent.tool_content import (
    decode_tool_content,
    encode_tool_content,
    tool_content_attachments,
)
from dlightrag.engine.agent.tools.files import (
    PreparedImageAttachment,
    ViewArgs,
    read_tool,
    view_tool,
)
from dlightrag.engine.ai.media import decode_image_base64
from dlightrag.engine.answer.resources.converters import ResourceConversionError
from dlightrag.engine.answer.resources.models import (
    ResourceInput,
    ResourceRegistryError,
    TextWindowBudget,
)
from dlightrag.engine.answer.resources.registry import ResourceRegistry
from dlightrag.engine.answer.resources.snapshots import ConversionSnapshot
from dlightrag.engine.answer.tools.resources import make_resource_reader, make_resource_viewer
from tests.tool_helpers import tool_runtime
from tests.unit.conftest import answer_image_policy
from tests.unit.test_resource_visual import pdf_bytes


def png():
    buffer = io.BytesIO()
    Image.new("RGB", (24, 24), (20, 10, 0)).save(buffer, "PNG")
    return buffer.getvalue()


def preparer(max_images=8):
    budget = answer_image_policy(max_images=max_images).new_budget()

    def prepare(data, label):
        block = budget.add_base64(base64.b64encode(data).decode(), label=label)
        if block is None:
            return None
        content, media = decode_image_base64(block["image_url"]["url"])
        return PreparedImageAttachment(content, media or "image/png", content != data)

    return prepare


def tools(registry, *, max_images=8, environment=None):
    access = AccessScheduler()
    return (
        read_tool(
            environment,
            access,
            resource_reader=make_resource_reader(registry, TextWindowBudget(1000)),
        ),
        view_tool(
            environment,
            access,
            resource_viewer=make_resource_viewer(registry),
            image_preparer=preparer(max_images),
        ),
    )


async def call(tool, **args):
    return await tool.execute(
        tool.input_model.model_validate(args), tool_runtime(tool_name=tool.name)
    )


@pytest.mark.parametrize(
    "args",
    [
        {},
        {"path": "a", "url": "https://example.com"},
        {"path": "a", "locator": "1"},
        {"resource_id": "res-a", "focus": "x"},
        {"url": "https://example.com", "cursor": "x"},
        {"resource_id": "res-a", "http": {}},
        {"resource_id": "res-a", "locator": "1", "cursor": "x"},
    ],
)
def test_view_rejects_ambiguous_or_legacy_arguments(args):
    with pytest.raises(ValidationError):
        ViewArgs.model_validate(args)


async def test_read_image_returns_guidance_only_and_view_attaches_located_pixels():
    async with ResourceRegistry() as registry:
        resource = registry.register(ResourceInput(filename="plot.png", content=png()))
        read, view = tools(registry)
        text = await call(read, resource_id=resource)
        assert not tool_content_attachments(text.parts)
        assert "view(resource_id=" in text.text_content
        pixels = await call(view, resource_id=resource)
        (attachment,) = tool_content_attachments(pixels.parts)
        assert attachment.data == png()
        assert attachment.source is not None
        assert attachment.source.resource_id == resource
        assert attachment.source is not None
        assert attachment.source.kind == "image"
        restored = decode_tool_content(encode_tool_content(pixels.parts))
        (restored_attachment,) = tool_content_attachments(restored)
        assert restored_attachment.source == attachment.source
        assert not restored_attachment.data


async def test_workspace_read_image_is_text_and_view_rejects_escape_and_documents(tmp_path):
    (tmp_path / "a.png").write_bytes(png())
    (tmp_path / "doc.pdf").write_bytes(pdf_bytes())
    async with ResourceRegistry() as registry:
        read, view = tools(registry, environment=LocalExecutionEnvironment(tmp_path))
        assert not tool_content_attachments((await call(read, path="a.png")).parts)
        assert tool_content_attachments((await call(view, path="a.png")).parts)
        assert (await call(view, path="../escape.png")).is_error
        assert (await call(view, path="doc.pdf")).is_error


async def test_pdf_view_bypasses_failed_text_extraction(monkeypatch):
    async def fail(*args, **kwargs):
        raise ResourceConversionError("ordinary parser failure")

    monkeypatch.setattr("dlightrag.engine.answer.resources.registry.convert_resource", fail)
    async with ResourceRegistry() as registry:
        resource = registry.register(ResourceInput(filename="paper.pdf", content=pdf_bytes()))
        read, view = tools(registry)
        text = await call(read, resource_id=resource)
        assert "conversion_failed" in text.text_content
        assert "Physical PDF page count: 3" in text.text_content
        result = await call(view, resource_id=resource, locator="2")
        (attachment,) = tool_content_attachments(result.parts)
        assert attachment.source is not None
        assert attachment.source.page == 2
        assert attachment.source is not None
        assert not attachment.source.overview


async def test_pdf_overview_actual_coverage_aggregate_budget_and_signed_recovery_cursor():
    data = pdf_bytes(3)
    async with ResourceRegistry(resource_secret=b"r", cursor_secret=b"c") as registry:
        resource = registry.register(ResourceInput(filename="paper.pdf", content=data))
        _, view = tools(registry, max_images=1)
        result = await call(view, resource_id=resource)
        assert "physical pages 1-1 of 3 only" in result.text_content
        cursor = result.protected_text.split("cursor='")[1].split("'")[0]
        with pytest.raises(ResourceRegistryError, match="remaining model image budget"):
            await call(view, resource_id=resource, cursor=cursor)
    async with ResourceRegistry(resource_secret=b"r", cursor_secret=b"c") as recovered:
        assert recovered.register(ResourceInput(filename="paper.pdf", content=data)) == resource
        _, view = tools(recovered, max_images=1)
        second = await call(view, resource_id=resource, cursor=cursor)
        (attachment,) = tool_content_attachments(second.parts)
        assert attachment.source is not None
        assert attachment.source.page == 2
        with pytest.raises(ResourceRegistryError):
            await call(view, resource_id=resource, cursor=cursor + "x")


def docx_images(count):
    doc = Document()
    doc.add_paragraph("Revenue was 123.")
    for _ in range(count):
        doc.add_picture(io.BytesIO(png()))
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


async def test_duplicate_occurrences_membership_inventory_and_snapshot_reuse(monkeypatch):
    data = docx_images(12)
    async with ResourceRegistry(resource_secret=b"r", cursor_secret=b"c") as registry:
        resource = registry.register(ResourceInput(filename="a.docx", content=data))
        other = registry.register(ResourceInput(filename="b.docx", content=docx_images(1)))
        read, view = tools(registry)
        result = await call(read, resource_id=resource)
        assert "more: read(" in result.text_content
        assets = [
            a for a in result.effects.attached_resources if a.resource_kind == "conversion_asset"
        ]
        assert len(assets) == 12
        assert len({a.resource_id for a in assets}) == 12
        assert len({a.content for a in assets}) == 1
        with pytest.raises(ResourceRegistryError, match="unknown visual handle"):
            await call(view, resource_id=other, locator=assets[0].resource_id)
        cursor = result.text_content.split("more: read(")[1].split("cursor='")[1].split("'")[0]
        page = await call(read, resource_id=resource, cursor=cursor)
        assert "Visual inventory" in page.text_content
        stored = {a.resource_id: a.content for a in result.effects.attached_resources}
        snapshot = ConversionSnapshot.restore(stored[f"{resource}-conversion"], stored)

    async def forbidden(*args, **kwargs):
        raise AssertionError("adopted snapshots never reparse")

    monkeypatch.setattr("dlightrag.engine.answer.resources.registry.convert_resource", forbidden)
    async with ResourceRegistry(resource_secret=b"r", cursor_secret=b"c") as registry:
        assert registry.register(ResourceInput(filename="a.docx", content=data)) == resource
        registry.adopt_conversion_snapshot(snapshot)
        read, view = tools(registry)
        assert "Revenue was 123" in (await call(read, resource_id=resource)).text_content
        pixels = await call(view, resource_id=resource, locator=assets[-1].resource_id)
        source = tool_content_attachments(pixels.parts)[0].source
        assert source is not None
        assert source.handle_id == assets[-1].resource_id


async def test_extensionless_url_image_is_classified_after_acquisition_and_reuses_snapshot(
    monkeypatch,
):
    from types import SimpleNamespace

    calls = []

    async def fetch(url, **kwargs):
        calls.append(url)
        return SimpleNamespace(content=png(), media_type="image/png", final_url=url)

    monkeypatch.setattr("dlightrag.engine.answer.resources.registry.fetch_public_http", fetch)
    async with ResourceRegistry() as registry:
        read, view = tools(registry)
        result = await call(view, url="https://example.com/asset")
        (attachment,) = tool_content_attachments(result.parts)
        assert attachment.source is not None
        resource = attachment.source.resource_id
        assert attachment.source is not None
        assert attachment.source.kind == "image"
        assert not tool_content_attachments((await call(read, resource_id=resource)).parts)
        await call(view, resource_id=resource)
        with pytest.raises(ResourceRegistryError, match="cannot replace"):
            await call(view, url="https://example.com/asset", http={"accept": "image/webp"})
        assert len(calls) == 1


async def test_conversion_cancellation_does_not_overlap_native_work_or_cleanup(monkeypatch):
    import asyncio
    import threading

    started, release = threading.Event(), threading.Event()
    calls = []

    def native(*args, **kwargs):
        calls.append(1)
        started.set()
        assert release.wait(5)
        return "adopted"

    monkeypatch.setattr("anydoc.to_markdown_bytes", native)
    registry = ResourceRegistry()
    resource = registry.register(ResourceInput(filename="a.pdf", content=pdf_bytes(1)))
    first = asyncio.create_task(registry.read(resource, max_window_tokens=1000))
    assert await asyncio.to_thread(started.wait, 5)
    first.cancel()
    with pytest.raises(asyncio.CancelledError):
        await first
    second = asyncio.create_task(registry.read(resource, max_window_tokens=1000))
    await asyncio.sleep(0)
    release.set()
    from dlightrag.engine.answer.resources.converters import ConversionLimitError

    with pytest.raises(ConversionLimitError):
        await second
    assert len(calls) == 1
    await registry.aclose()
