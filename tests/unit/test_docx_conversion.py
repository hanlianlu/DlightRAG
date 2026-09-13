# Copyright 2025-2026 Hanlian Lu. SPDX-License-Identifier: Apache-2.0
"""Actual unified DOCX adapter gold, occurrence safety, fallback and recovery."""

import asyncio
import hashlib
import io
import json
import threading
import zipfile
from types import SimpleNamespace
from typing import cast

import anydoc
import pytest
from docx import Document

from dlightrag.engine.answer.resources import converters
from dlightrag.engine.answer.resources.converters import ConversionLimitError, convert_resource
from dlightrag.engine.answer.resources.docx_assets import AssetBindingError, docx_asset_occurrences
from dlightrag.engine.answer.resources.models import ResourceInput
from dlightrag.engine.answer.resources.registry import ResourceRegistry
from dlightrag.engine.answer.resources.snapshots import ConversionSnapshot
from scripts.anydoc_pilot import evaluate
from scripts.anydoc_pilot_fixtures import _zip_replace, generate


@pytest.fixture(scope="module")
def fixtures(tmp_path_factory):
    root = tmp_path_factory.mktemp("unified-docx")
    return root, {g["name"]: g for g in generate(root)}


@pytest.mark.parametrize(
    "name",
    [
        "docx-text.docx",
        "docx-number-table.docx",
        "docx-footnote.docx",
        "docx-empty.docx",
        "docx-repeat-image.docx",
    ],
)
async def test_unified_docx_gold_and_adopted_inventory(fixtures, name, monkeypatch):
    root, gold = fixtures
    data = (root / name).read_bytes()
    calls = []
    markdown, structured = anydoc.to_markdown_bytes, anydoc.to_document

    def md(*args, **kwargs):
        calls.append("markdown")
        assert kwargs == {"format": "docx", "ocr": "reject"}
        return markdown(*args, **kwargs)

    def doc(*args, **kwargs):
        calls.append("document")
        return structured(*args, **kwargs)

    monkeypatch.setattr(anydoc, "to_markdown_bytes", md)
    monkeypatch.setattr(anydoc, "to_document", doc)
    async with ResourceRegistry(resource_secret=b"gold", cursor_secret=b"cursor") as registry:
        resource = registry.register(ResourceInput(filename=name, content=data))
        first = await registry.read(resource, max_window_tokens=4000)
        assert first == await registry.read(resource, max_window_tokens=4000)
        effects = registry.conversion_effects(resource)
        snapshot = ConversionSnapshot.restore(
            effects[-1].content, {e.resource_id: e.content for e in effects[:-1]}
        )
        assert snapshot.converter == "firecrawl-anydoc"
        assert snapshot.converter_version == "0.2.4"
        assert snapshot.input_digest == hashlib.sha256(data).hexdigest()
        assert snapshot.fallback_reason is None
        assert snapshot.extraction_status == (
            "no_extracted_text" if name == "docx-empty.docx" else "usable_text_unverified_coverage"
        )
        # Gold requires discoverable occurrences, not fabricated inline links.
        result = {
            "text": snapshot.text,
            "assets": [
                {"sha256": hashlib.sha256(v.data).hexdigest(), "anchor": v.anchor}
                for v in snapshot.visuals
            ],
            "visual_references": len(first.visual_handles),
        }
        assert evaluate(gold[name], result)["passed"]
        for visual in snapshot.visuals:
            assert visual.anchor is None and visual.origin_part == "word/media/image1.png"
            assert visual.handle_id not in snapshot.text
            assert (await registry.visual_asset(resource, visual.handle_id)).data == visual.data
        if snapshot.visuals:
            assert len({v.handle_id for v in snapshot.visuals}) == 2
            assert snapshot.visuals[0].data == snapshot.visuals[1].data
    async with ResourceRegistry(resource_secret=b"gold", cursor_secret=b"cursor") as restored:
        restored.register(ResourceInput(filename=name, content=data))
        restored.adopt_conversion_snapshot(snapshot)
        assert await restored.read(resource, max_window_tokens=4000) == first
    assert calls == ["markdown", "document"]


@pytest.mark.parametrize(
    "error", [anydoc.MalformedError, anydoc.MissingPartError, ImportError, OSError]
)
async def test_single_ordinary_or_init_fallback(fixtures, monkeypatch, error):
    root, _ = fixtures
    calls = []
    incumbent = converters._convert_markitdown

    def fail(*args, **kwargs):
        calls.append("candidate-ended")
        raise error("synthetic")

    def fallback(*args):
        assert calls == ["candidate-ended"]
        calls.append("fallback")
        return incumbent(*args)

    monkeypatch.setattr(
        converters if error in {ImportError, OSError} else anydoc,
        "_load_anydoc" if error in {ImportError, OSError} else "to_markdown_bytes",
        fail,
    )
    monkeypatch.setattr(converters, "_convert_markitdown", fallback)
    result = await convert_resource(
        (root / "docx-text.docx").read_bytes(), filename="a.docx", declared_mime=None
    )
    assert "718.40" in result.text
    assert result.converter == "markitdown"
    assert result.fallback_reason and error.__name__ in result.fallback_reason
    assert calls == ["candidate-ended", "fallback"]


@pytest.mark.parametrize(
    "error",
    [
        anydoc.NeedsOcrError,
        anydoc.ResourceLimitError,
        anydoc.UnsupportedError,
        anydoc.EncryptedError,
        MemoryError,
        RuntimeError,
        ValueError,
    ],
)
async def test_candidate_typed_terminals_never_fallback(fixtures, monkeypatch, error):
    root, _ = fixtures
    exc = error("synthetic")
    if isinstance(exc, anydoc.NeedsOcrError):
        exc.pages, exc.page_count = [2], 3

    def fail(*args, **kwargs):
        raise exc

    def forbidden(*args):
        pytest.fail("terminal must not invoke incumbent")

    monkeypatch.setattr(anydoc, "to_markdown_bytes", fail)
    monkeypatch.setattr(converters, "_convert_markitdown", forbidden)
    source = ResourceInput(filename="a.docx", content=(root / "docx-text.docx").read_bytes())
    async with ResourceRegistry() as registry:
        resource = registry.register(source)
        if error in {anydoc.NeedsOcrError, anydoc.UnsupportedError, anydoc.EncryptedError}:
            result = await registry.read(resource, max_window_tokens=1000)
            assert result.content == ""
            raw = json.loads(registry.conversion_effects(resource)[-1].content)
            assert raw["converter"] == "firecrawl-anydoc"
            assert raw["fallback_reason"] is None
            if error == anydoc.NeedsOcrError:
                assert raw["known_ocr_pages"] == [2] and raw["known_page_count"] == 3
                assert "Known OCR pages: [2] of 3" in (result.note or "")
                assert result.extraction_status == "known_incomplete"
        else:
            with pytest.raises((ConversionLimitError, MemoryError, RuntimeError, ValueError)):
                await registry.read(resource, max_window_tokens=1000)


async def test_native_timeout_cannot_start_late_fallback(fixtures, monkeypatch):
    root, _ = fixtures
    started, release = threading.Event(), threading.Event()
    calls = []

    def native(*args, **kwargs):
        calls.append("candidate")
        started.set()
        assert release.wait(5)
        raise anydoc.MalformedError("ended only after cancellation")

    def forbidden(*args):
        pytest.fail("late fallback")

    monkeypatch.setattr(anydoc, "to_markdown_bytes", native)
    monkeypatch.setattr(converters, "_convert_markitdown", forbidden)
    registry = ResourceRegistry()
    resource = registry.register(
        ResourceInput(filename="a.docx", content=(root / "docx-text.docx").read_bytes())
    )
    read = asyncio.create_task(registry.read(resource, max_window_tokens=1000))
    assert await asyncio.to_thread(started.wait, 5)
    read.cancel()
    with pytest.raises(asyncio.CancelledError):
        await read
    close = asyncio.create_task(registry.aclose())
    await asyncio.sleep(0)
    assert not close.done()
    release.set()
    await close
    assert calls == ["candidate"]


async def test_shared_deadline_exhaustion_never_retries(fixtures, monkeypatch):
    root, _ = fixtures
    now = [10.0]
    monkeypatch.setattr(converters.time, "monotonic", lambda: now[0])

    def fail(*args, **kwargs):
        now[0] += 121
        raise anydoc.MalformedError("late")

    monkeypatch.setattr(anydoc, "to_markdown_bytes", fail)
    monkeypatch.setattr(
        converters, "_convert_markitdown", lambda *args: pytest.fail("renewed budget")
    )
    with pytest.raises(ConversionLimitError):
        await convert_resource(
            (root / "docx-text.docx").read_bytes(), filename="a.docx", declared_mime=None
        )


@pytest.mark.parametrize("mutation", ["external", "missing", "header"])
async def test_unmapped_source_visuals_are_known_incomplete(fixtures, mutation):
    root, _ = fixtures
    data = (root / "docx-repeat-image.docx").read_bytes()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        rels = archive.read("word/_rels/document.xml.rels")
    if mutation == "external":
        data = _zip_replace(
            data,
            {
                "word/_rels/document.xml.rels": rels.replace(
                    b'Target="media/image1.png"',
                    b'Target="https://invalid.example/image.png" TargetMode="External"',
                )
            },
        )
    elif mutation == "missing":
        data = _zip_replace(data, {"word/media/image1.png": None})
    else:
        doc = Document(io.BytesIO((root / "docx-text.docx").read_bytes()))
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            doc.sections[0].header.paragraphs[0].add_run().add_picture(
                io.BytesIO(archive.read("word/media/image1.png"))
            )
        buf = io.BytesIO()
        doc.save(buf)
        data = buf.getvalue()
    result = await convert_resource(data, filename="a.docx", declared_mime=None)
    assert result.converter == "firecrawl-anydoc"
    assert result.extraction_status == "known_incomplete"
    assert result.fallback_reason is None
    assert result.note


@pytest.mark.parametrize(
    "field,value",
    [
        ("origin_part", "../foreign.png"),
        ("origin_part", "word/media/missing.png"),
        ("data", b"foreign"),
        ("media_type", "image/jpeg"),
    ],
)
def test_candidate_assets_must_match_exact_source(fixtures, field, value):
    root, _ = fixtures
    data = (root / "docx-repeat-image.docx").read_bytes()
    document = anydoc.to_document(data, format="docx")
    asset = document.assets[0]
    forged = {name: getattr(asset, name) for name in ("id", "origin_part", "data", "media_type")}
    forged[field] = value
    fake = SimpleNamespace(
        assets=[SimpleNamespace(**forged)], blocks=document.blocks, notes=document.notes
    )
    with pytest.raises(AssetBindingError):
        docx_asset_occurrences(cast(anydoc.Document, fake), data)


async def test_unsafe_source_invokes_neither_parser(fixtures, monkeypatch):
    root, _ = fixtures
    monkeypatch.setattr(
        anydoc, "to_markdown_bytes", lambda *a, **k: pytest.fail("unsafe candidate")
    )
    monkeypatch.setattr(
        converters, "_convert_markitdown", lambda *a: pytest.fail("unsafe fallback")
    )
    with pytest.raises(converters.UnsafeArchiveError):
        await convert_resource(
            (root / "docx-unsafe.docx").read_bytes(), filename="a.docx", declared_mime=None
        )


async def test_table_image_origin_cell_is_one_occurrence(fixtures):
    root, _ = fixtures
    with zipfile.ZipFile(root / "docx-repeat-image.docx") as archive:
        image = archive.read("word/media/image1.png")
    document = Document()
    table = document.add_table(rows=1, cols=2)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.paragraphs[0].add_run().add_picture(io.BytesIO(image))
    stream = io.BytesIO()
    document.save(stream)
    result = await convert_resource(stream.getvalue(), filename="table.docx", declared_mime=None)
    assert result.extraction_status != "known_incomplete"
    assert len(result.visuals) == 1 and result.visuals[0].data == image


async def test_rich_docx_fallback_preserves_occurrences_and_membership(fixtures, monkeypatch):
    root, _ = fixtures

    def fail(*args, **kwargs):
        raise anydoc.MalformedError("ordinary")

    monkeypatch.setattr(anydoc, "to_markdown_bytes", fail)
    result = await convert_resource(
        (root / "docx-repeat-image.docx").read_bytes(), filename="rich.docx", declared_mime=None
    )
    assert result.converter == "markitdown" and result.fallback_reason
    assert result.extraction_status == "usable_text_unverified_coverage"
    assert len(result.visuals) == 2
    assert len({v.handle_id for v in result.visuals}) == 2
    assert all(v.origin_part == "word/media/image1.png" for v in result.visuals)


async def test_failed_fallback_metadata_is_adopted_once(fixtures, monkeypatch):
    root, _ = fixtures
    calls = []

    def candidate(*args, **kwargs):
        calls.append("candidate")
        raise anydoc.MissingPartError("ordinary")

    def incumbent(*args):
        calls.append("incumbent")
        raise converters.ResourceConversionError("also failed")

    monkeypatch.setattr(anydoc, "to_markdown_bytes", candidate)
    monkeypatch.setattr(converters, "_convert_markitdown", incumbent)
    async with ResourceRegistry() as registry:
        resource = registry.register(
            ResourceInput(filename="a.docx", content=(root / "docx-text.docx").read_bytes())
        )
        first = await registry.read(resource, max_window_tokens=1000)
        assert (
            (await registry.read(resource, max_window_tokens=1000)).extraction_status
            == first.extraction_status
            == "conversion_failed"
        )
        snapshot = json.loads(registry.conversion_effects(resource)[-1].content)
        assert snapshot["converter"] == "markitdown"
        assert "MissingPartError" in snapshot["fallback_reason"]
        assert calls == ["candidate", "incumbent"]


def test_unknown_structured_asset_is_not_silently_lost(fixtures):
    root, _ = fixtures
    data = (root / "docx-repeat-image.docx").read_bytes()
    document = anydoc.to_document(data, format="docx")
    fake = SimpleNamespace(assets=[], blocks=document.blocks, notes=document.notes)
    visuals, note = docx_asset_occurrences(cast(anydoc.Document, fake), data)
    assert visuals == [] and note and "incomplete" in note


async def test_real_native_depth_limit_is_terminal(fixtures, monkeypatch):
    root, _ = fixtures
    data = (root / "docx-text.docx").read_bytes()
    deep = (
        b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        + b"<w:body>" * 260
        + b"</w:body>" * 260
        + b"</w:document>"
    )
    data = _zip_replace(data, {"word/document.xml": deep})
    monkeypatch.setattr(
        converters, "_convert_markitdown", lambda *a: pytest.fail("native limit fallback")
    )
    with pytest.raises(ConversionLimitError) as error:
        await convert_resource(data, filename="depth.docx", declared_mime=None)
    assert isinstance(error.value.__cause__, anydoc.ResourceLimitError)


@pytest.mark.parametrize("failure", [anydoc.ResourceLimitError, MemoryError])
async def test_resource_refusal_tool_effects_restore_without_reparse(
    fixtures, monkeypatch, failure
):
    from tests.unit.test_resource_tools import call, tools

    root, _ = fixtures

    def fail(*args, **kwargs):
        raise failure("terminal")

    monkeypatch.setattr(anydoc, "to_markdown_bytes", fail)
    source = ResourceInput(filename="a.docx", content=(root / "docx-text.docx").read_bytes())
    async with ResourceRegistry(resource_secret=b"safety") as registry:
        resource = registry.register(source)
        read, _ = tools(registry)
        result = await call(read, resource_id=resource)
        assert result.is_error and "safety_refused" in result.text_content
        effect = result.effects.attached_resources[-1]
        snapshot = ConversionSnapshot.restore(effect.content, {})
        assert snapshot.extraction_status == "safety_refused"
    monkeypatch.setattr(
        anydoc, "to_markdown_bytes", lambda *a, **k: pytest.fail("refused snapshot reparsed")
    )
    async with ResourceRegistry(resource_secret=b"safety") as restored:
        restored.register(source)
        restored.adopt_conversion_snapshot(snapshot)
        read, _ = tools(restored)
        assert (await call(read, resource_id=resource)).is_error
