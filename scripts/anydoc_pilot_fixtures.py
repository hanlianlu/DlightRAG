# Copyright 2025-2026 Hanlian Lu. SPDX-License-Identifier: Apache-2.0
"""Small synthetic semantic gold; no incumbent output is used to create expectations."""

from __future__ import annotations

import datetime
import hashlib
import io
import json
import re
import zipfile
from pathlib import Path
from typing import Any, cast


def _zip_replace(data: bytes, changes: dict[str, bytes | None]) -> bytes:
    target = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(target, "w") as dest:
        for name in source.namelist():
            value = changes.get(name, source.read(name))
            if value is not None:
                dest.writestr(name, value)
        for name, value in changes.items():
            if name not in source.namelist() and value is not None:
                dest.writestr(name, value)
    return target.getvalue()


def _pdf(pages: list[tuple[list[str], int]]) -> bytes:
    """Minimal real PDF with Helvetica text and optional repeated raster regions.

    Raster pixels contain SCAN ONLY 83.50, generated independently of PDF text.
    No OCR text layer or alt text is added to the image object.
    """
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (400, 100), "white")
    ImageDraw.Draw(image).text((10, 30), "SCAN ONLY 83.50", fill="black", font_size=30)
    pixels = image.tobytes()
    objects = [b"", b"", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    if any(images for _, images in pages):
        objects.append(
            b"<< /Type /XObject /Subtype /Image /Width 400 /Height 100 /ColorSpace /DeviceRGB "
            b"/BitsPerComponent 8 /Length "
            + str(len(pixels)).encode()
            + b" >>\nstream\n"
            + pixels
            + b"\nendstream"
        )
    page_ids = []
    for lines, images in pages:
        commands = ["BT /F1 16 Tf 50 750 Td"]
        for line in lines:
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            commands.append(f"({escaped}) Tj 0 -28 Td")
        commands.append("ET")
        for index in range(images):
            commands.append(f"q 400 0 0 100 50 {350 - index * 120} cm /Im1 Do Q")
        stream = "\n".join(commands).encode("ascii")
        page_id = len(objects) + 1
        page_ids.append(page_id)
        image_resource = "/XObject << /Im1 4 0 R >>" if images else ""
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources "
            f"<< /Font << /F1 3 0 R >> {image_resource} >> "
            f"/Contents {page_id + 1} 0 R >>".encode()
        )
        objects.append(
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        )
    objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[1] = (
        f"<< /Type /Pages /Count {len(pages)} /Kids [{' '.join(f'{p} 0 R' for p in page_ids)}] >>".encode()
    )
    result = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(result))
        result.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(result)
    result.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        result.extend(f"{offset:010d} 00000 n \n".encode())
    result.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(result)


def generate(root: Path) -> list[dict[str, Any]]:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches as DocInches
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as SheetImage
    from PIL import Image, ImageDraw
    from pptx import Presentation
    from pptx.shapes.placeholder import SlidePlaceholder
    from pptx.util import Inches

    root.mkdir(parents=True, exist_ok=True)
    gold: list[dict[str, Any]] = []
    image = Image.new("RGB", (160, 80), "white")
    ImageDraw.Draw(image).text((8, 25), "FIGURE 47", fill="black", font_size=20)
    stream = io.BytesIO()
    image.save(stream, "PNG")
    png = stream.getvalue()
    image_digest = hashlib.sha256(png).hexdigest()

    def add(name: str, data: bytes, facts: list[str], **extra: Any) -> None:
        if (
            Path(name).suffix in {".docx", ".pptx", ".xlsx"}
            and extra.get("expected") != "host_safety_refusal"
        ):
            normalized = io.BytesIO()
            with (
                zipfile.ZipFile(io.BytesIO(data)) as source,
                zipfile.ZipFile(normalized, "w") as target,
            ):
                for part in sorted(source.namelist()):
                    content = source.read(part)
                    if part == "docProps/core.xml":
                        content = re.sub(
                            rb"(<dcterms:(?:created|modified)[^>]*>)[^<]+",
                            rb"\g<1>2020-01-01T00:00:00Z",
                            content,
                        )
                    info = zipfile.ZipInfo(part, date_time=(2020, 1, 1, 0, 0, 0))
                    info.compress_type = zipfile.ZIP_DEFLATED
                    target.writestr(info, content)
            data = normalized.getvalue()
        (root / name).write_bytes(data)
        gold.append(
            {
                "name": name,
                "sha256": hashlib.sha256(data).hexdigest(),
                "bytes": len(data),
                "facts": facts,
                "patterns": [],
                "expected": "text",
                "asset_occurrences": 0,
                "anchors": [],
                **extra,
            }
        )

    add(
        "pdf-text.pdf",
        _pdf(
            [
                (["ALPHA revenue 127.25", "Physical page one"], 0),
                (["BETA count 42", "Physical page two"], 0),
            ]
        ),
        ["ALPHA", "127.25", "BETA", "42", "Physical page one", "Physical page two"],
        pages=2,
    )
    add(
        "pdf-table.pdf",
        _pdf(
            [
                (
                    [
                        "Inventory",
                        "Item       Units       Price",
                        "COPPER       12       19.75",
                        "TIN       8       3.50",
                        "Total 265.00",
                    ],
                    0,
                )
            ]
        ),
        ["COPPER", "12", "19.75", "TIN", "8", "3.50", "265.00"],
        patterns=[r"COPPER\s+12\s+19\.75", r"TIN\s+8\s+3\.50"],
        pages=1,
    )
    add(
        "pdf-scan.pdf",
        _pdf([([], 1)]),
        ["SCAN ONLY", "83.50"],
        expected="ocr",
        pages=1,
        ocr_pages=[1],
    )
    add(
        "pdf-mixed.pdf",
        _pdf([(["COVER total 52"], 0), ([], 1)]),
        ["COVER", "52", "SCAN ONLY", "83.50"],
        expected="ocr",
        pages=2,
        ocr_pages=[2],
    )
    add(
        "pdf-text-image.pdf",
        _pdf([(["TEXT REGION total 52"], 1)]),
        ["TEXT REGION", "52", "SCAN ONLY", "83.50"],
        expected="incomplete",
        pages=1,
    )
    add(
        "pdf-repeat-image.pdf",
        _pdf([(["FIGURE references A and B"], 2)]),
        ["FIGURE references A and B", "SCAN ONLY", "83.50"],
        expected="incomplete",
        pages=1,
        image_occurrences=2,
    )

    def doc_bytes(doc: Any) -> bytes:
        out = io.BytesIO()
        doc.save(out)
        return out.getvalue()

    doc = Document()
    doc.add_heading("Contract Delta", 0)
    doc.add_paragraph("Delivery 2026-08-19. Amount USD 718.40. Quantity 23.")
    add("docx-text.docx", doc_bytes(doc), ["Contract Delta", "2026-08-19", "718.40", "23"])
    doc = Document()
    doc.add_heading("Assembly", 1)
    doc.add_paragraph("Install bolt 14", "List Number")
    doc.add_paragraph("Apply torque 27 Nm", "List Number")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Part", "Cost"
    row = table.add_row().cells
    row[0].text, row[1].text = "AXLE", "62.90"
    add(
        "docx-number-table.docx",
        doc_bytes(doc),
        ["Install bolt 14", "Apply torque 27 Nm", "AXLE", "62.90"],
        patterns=[r"1\.\s+Install bolt 14", r"2\.\s+Apply torque 27 Nm", r"AXLE\s*\|\s*62\.90"],
    )
    doc = Document()
    para = doc.add_paragraph("Footnote reference")
    ref = OxmlElement("w:footnoteReference")
    ref.set(qn("w:id"), "1")
    para.add_run()._r.append(ref)
    data = doc_bytes(doc)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        rels = archive.read("word/_rels/document.xml.rels").replace(
            b"</Relationships>",
            b'<Relationship Id="rIdPilotFootnote" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>',
        )
        types = archive.read("[Content_Types].xml").replace(
            b"</Types>",
            b'<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>',
        )
    footnotes = b'<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="1"><w:p><w:r><w:t>Warranty 36 months</w:t></w:r></w:p></w:footnote></w:footnotes>'
    add(
        "docx-footnote.docx",
        _zip_replace(
            data,
            {
                "word/footnotes.xml": footnotes,
                "word/_rels/document.xml.rels": rels,
                "[Content_Types].xml": types,
            },
        ),
        ["Footnote reference", "Warranty 36 months"],
    )
    doc = Document()
    doc.add_paragraph("Repeated figures with separate occurrences")
    for _ in range(2):
        doc.add_picture(io.BytesIO(png), width=DocInches(1))
    add(
        "docx-repeat-image.docx",
        doc_bytes(doc),
        ["Repeated figures with separate occurrences"],
        asset_occurrences=2,
        asset_digest=image_digest,
        expected="rich",
    )

    deck = Presentation()
    for title, body, note in [
        ("North slide 1", "Units 38", "Speaker margin 11.25"),
        ("South slide 2", "Units 49", "Speaker margin 21.75"),
    ]:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        title_shape = slide.shapes.title
        notes = slide.notes_slide.notes_text_frame
        if title_shape is None or notes is None:
            raise RuntimeError("synthetic slide layout is missing title/notes")
        title_shape.text = title
        cast(SlidePlaceholder, slide.placeholders[1]).text = body
        notes.text = note
    deck_data = doc_bytes(deck)
    add(
        "pptx-slides-notes.pptx",
        deck_data,
        ["North slide 1", "Units 38", "11.25", "South slide 2", "Units 49", "21.75"],
        slides=2,
    )
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(5), Inches(2)).table
    for cell, text in [
        (table.cell(0, 0), "ZONE"),
        (table.cell(0, 1), "LOAD"),
        (table.cell(1, 0), "WEST"),
        (table.cell(1, 1), "88.50"),
    ]:
        cell.text = text
    add(
        "pptx-table.pptx",
        doc_bytes(deck),
        ["ZONE", "LOAD", "WEST", "88.50"],
        patterns=[r"WEST\s*\|\s*88\.50"],
        slides=1,
    )
    for index in range(2):
        slide.shapes.add_picture(io.BytesIO(png), Inches(index + 1), Inches(4), width=Inches(1))
    add(
        "pptx-repeat-image.pptx",
        doc_bytes(deck),
        ["WEST", "88.50"],
        asset_occurrences=2,
        asset_digest=image_digest,
        expected="rich",
        slides=1,
    )
    add(
        "pptx-missing-slide.pptx",
        _zip_replace(deck_data, {"ppt/slides/slide2.xml": None}),
        ["North slide 1", "South slide 2"],
        expected="malformed",
        slides=2,
    )

    book = Workbook()
    book.properties.created = datetime.datetime(2020, 1, 1)
    book.properties.modified = datetime.datetime(2020, 1, 1)
    sheet = book.active
    if sheet is None:
        raise RuntimeError("synthetic workbook is missing its initial sheet")
    sheet.title = "Metrics"
    sheet.append(["Ratio", "Date", "Label"])
    sheet.append([0.125, datetime.date(2026, 8, 19), "A|B"])
    sheet["A2"].number_format = "0.0%"
    sheet["B2"].number_format = "yyyy-mm-dd"
    sheet.merge_cells("A4:B4")
    sheet["A4"] = "Merged total 74"
    sheet.add_image(SheetImage(io.BytesIO(png)), "D5")
    add(
        "xlsx-display-image.xlsx",
        doc_bytes(book),
        ["12.5%", "2026-08-19", "A|B", "Merged total 74"],
        asset_occurrences=1,
        asset_digest=image_digest,
        anchors=["Metrics!D5"],
        expected="control",
    )
    add(
        "html-control.html",
        b'<html><body><h1>Harbor</h1><table><tr><td>Dock</td><td>9.75</td></tr></table><img src="https://invalid.example/pixel.png" alt="external figure"></body></html>',
        ["Harbor", "Dock", "9.75"],
        expected="control",
    )
    add(
        "csv-control.csv",
        b'Name,Amount,Note\r\n"East, branch",17.80,"line one\nline two"\r\n',
        ["East, branch", "17.80", "line one", "line two"],
        expected="control",
    )
    add("docx-empty.docx", doc_bytes(Document()), [], expected="empty")
    add("pdf-malformed.pdf", b"%PDF-1.4\nnot an object\n%%EOF", [], expected="malformed")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            zipfile.ZipInfo("word/document.xml", date_time=(2020, 1, 1, 0, 0, 0)),
            b"0" * 4096,
            compress_type=zipfile.ZIP_DEFLATED,
        )
    add("docx-unsafe.docx", out.getvalue(), [], expected="host_safety_refusal")
    (root / "gold.json").write_text(json.dumps(gold, indent=2) + "\n")
    return gold
