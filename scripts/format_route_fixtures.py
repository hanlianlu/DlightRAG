# Copyright 2025-2026 Hanlian Lu. SPDX-License-Identifier: Apache-2.0
"""Source-authored PDF/XLSX route gold, generated without network or licensed fonts.

Non-Latin font qualification is a separate optional local-only probe; no font
or resulting proprietary-font PDF is distributed with this generator.
"""

from __future__ import annotations

import datetime
import io
import re
import zipfile

from openpyxl import Workbook
from openpyxl.drawing.image import Image as SheetImage
from PIL import Image

from scripts.anydoc_pilot_fixtures import _pdf, _zip_replace

PDF_FACTS = ("DELTA revenue 111.25", "EPSILON count 22", "ZETA ratio 33.50")
XLSX_FACTS = (
    "Widget",
    "12.5%",
    "50.0%",
    "2026-08-19",
    "Merged note 74",
    "A\\|B",
    "28.5",
)
XLSX_ANCHORS = ("Q1!E7", "Q1!G9", "Q2!B6")


def pdf_fixture(kind: str) -> bytes:
    if kind == "multi":
        return _pdf([([fact, f"Physical page {i}"], 0) for i, fact in enumerate(PDF_FACTS, 1)])
    if kind == "scan":
        return _pdf([([], 1), ([], 1)])
    if kind == "mixed":
        return _pdf([(["PLAN total 52", "Q1 12", "Q2 40", "raster region below"], 1)])
    raise ValueError(f"unknown synthetic PDF {kind}")


def xlsx_fixture() -> tuple[bytes, bytes]:
    """Cached values are authored bytes, never calculated by the test or converter."""
    image = Image.new("RGB", (80, 40), (30, 80, 120))
    image.save(png := io.BytesIO(), "PNG")
    pixels = png.getvalue()
    book = Workbook()
    sheet = book.active
    if sheet is None:
        raise RuntimeError("generated workbook has no sheet")
    sheet.title = "Q1"
    sheet.append(["Product", "Ratio", "Date", "Flag"])
    sheet.append(["Widget", 0.125, datetime.date(2026, 8, 19), "A|B"])
    sheet.append(["Gadget", 0.5])
    sheet["B2"].number_format = sheet["B3"].number_format = "0.0%"
    sheet["C2"].number_format = "yyyy-mm-dd"
    sheet.merge_cells("A5:B5")
    sheet["A5"] = "Merged note 74"
    sheet.add_image(SheetImage(io.BytesIO(pixels)), "E7")
    sheet.add_image(SheetImage(io.BytesIO(pixels)), "G9")
    second = book.create_sheet("Q2")
    second.append(["Cached", "Uncached", "Currency", "Custom"])
    second.append(["=3*9.5", "=314159+271828", 1234.5, 7])
    second["C2"].number_format = '"$"#,##0.00'
    second["D2"].number_format = "0000"
    second.add_image(SheetImage(io.BytesIO(pixels)), "B6")
    book.save(stream := io.BytesIO())
    book.close()
    with zipfile.ZipFile(stream) as archive:
        xml = archive.read("xl/worksheets/sheet2.xml")
    xml, count = re.subn(rb'(<c r="A2"[^>]*><f>[^<]*</f>)<v></v>', rb"\g<1><v>28.5</v>", xml)
    if count != 1:
        raise RuntimeError("fixture cached value injection failed")
    return _zip_replace(stream.getvalue(), {"xl/worksheets/sheet2.xml": xml}), pixels
